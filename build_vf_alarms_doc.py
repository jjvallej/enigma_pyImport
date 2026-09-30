import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pathlib import Path

def set_cell(cell, text, font_name="Arial Narrow", font_size=Pt(10.5), bold=False, italic=False, bg_hex=None, text_color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.name = font_name
    if font_size:
        run.font.size = font_size
    run.bold = bold
    run.italic = italic
    if text_color:
        run.font.color.rgb = text_color
    if bg_hex:
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
        tcPr.append(shd)

def main():
    base_vf = Path("/home/jjvallej/work/enigma/pyimport/12. VF - DOCUMENTACION_CONEXIONES_AIRFLOW.docx")
    if not base_vf.exists():
        print(f"⚠️ Archivo plantilla {base_vf} no encontrado, usando docx básico.")
        doc = docx.Document()
    else:
        doc = docx.Document(str(base_vf))

    # 1. Update Table 0 (Metadata Table)
    if len(doc.tables) > 0:
        t0 = doc.tables[0]
        set_cell(t0.cell(1, 1), "Matriz de Eventos y Alarmas de DAGs de Apache Airflow")
        set_cell(t0.cell(2, 1), "Final")
        set_cell(t0.cell(5, 1), "10 de Septiembre de 2026")

    # 2. Update Table 1 (Control de Versiones)
    if len(doc.tables) > 1:
        t1 = doc.tables[1]
        set_cell(t1.cell(1, 0), "1.0.0")
        set_cell(t1.cell(1, 1), "31 de Agosto de 2026")
        set_cell(t1.cell(1, 2), "Matriz inicial de eventos de alarma (13 DAGs Medallion), códigos ALT-*, severidades y correo jhon.jairo.vallejo@gmail.com.")
        set_cell(t1.cell(1, 3), "Jhon Jairo Vallejo")
        
        if len(t1.rows) <= 2:
            t1.add_row()
        set_cell(t1.cell(2, 0), "Final")
        set_cell(t1.cell(2, 1), "10 de Septiembre de 2026")
        set_cell(t1.cell(2, 2), "Versión Final actualizada con la arquitectura completa de 15 DAGs de Apache Airflow, integrando GIS Espacial src_transform_spatial, Orquestador Maestro src_orchestrator_master y nuevos códigos ALT-TRF-004 y ALT-ORC-001.")
        set_cell(t1.cell(2, 3), "Jhon Jairo Vallejo")

    # 3. Title P19
    if len(doc.paragraphs) > 19:
        p19 = doc.paragraphs[19]
        p19.text = ""
        r = p19.add_run("MATRIZ DE EVENTOS Y ALARMAS DE DAGS DE APACHE AIRFLOW EN PRODUCCIÓN")
        r.font.name = "Arial"
        r.font.size = Pt(16)
        r.bold = True

    # 4. Table 2 - Alarms Matrix Table
    if len(doc.tables) > 2:
        t2 = doc.tables[2]

        alarms_data = [
            ["Código Alerta", "Capa / Etapa", "Descripción de la Falla / Evento", "Severidad", "DAGs Asociados", "Canal Notificación", "Procedimiento de Mitigación"],
            ["ALT-ING-001", "Ingesta (Raw)", "Timeout o fallo de conexión HTTP al descargar anexos SIPSA desde DANE", "CRÍTICO", "src_ingest_sipsa", "jhon.jairo.vallejo@gmail.com", "Reintento (1 retriable, 5 min), verificar portal dane.gov.co"],
            ["ALT-ING-002", "Ingesta (Raw)", "Fallo en la descarga del HTML del Índice ONI v5 desde NOAA CPC", "ALTO", "src_ingest_oni", "jhon.jairo.vallejo@gmail.com", "Reintento (1 retriable, 5 min), comprobar endpoint cpc.ncep.noaa.gov"],
            ["ALT-ING-003", "Ingesta (Raw)", "Error al descargar CSVs de cultivos Gobernación", "ALTO", "src_ingest_crops", "jhon.jairo.vallejo@gmail.com", "Revisar recurso CKAN Gobernación y conexión gobernacion_valle"],
            ["ALT-ING-004", "Ingesta (Raw)", "Falla al consultar catálogo de municipios en Datos Abiertos datos.gov.co", "MEDIO", "src_ingest_municipios", "jhon.jairo.vallejo@gmail.com", "Validar id iryd-wvq5 en datos.gov.co"],
            ["ALT-ING-005", "Ingesta (Raw)", "Fallo de conexión PostgreSQL/API a comentarios CKAN", "CRÍTICO", "src_ingest_ckan_comentarios", "jhon.jairo.vallejo@gmail.com", "Verificar credenciales de BD y endpoint CKAN"],
            ["ALT-ING-006", "Ingesta (Raw)", "Descarga incompleta o respuesta soft-404 HTML", "ALTO", "src_ingest_sipsa, src_ingest_crops", "jhon.jairo.vallejo@gmail.com", "Inspeccionar archivos en data/raw/ para detectar respuestas HTML ficticias"],
            ["ALT-LOD-001", "Carga (Bronze)", "Error de permisos IAM / ADC al escribir en Cloud Storage Bucket", "CRÍTICO", "Todos los DAGs src_load_*", "jhon.jairo.vallejo@gmail.com", "Verificar permisos de SA Composer sobre bucket gs://datalake_gdv_pdn"],
            ["ALT-LOD-002", "Carga (Bronze)", "Fallo al insertar registros o crear tablas staging en BigQuery Bronze", "CRÍTICO", "src_load_sipsa, src_load_oni, src_load_crops, src_load_municipios", "jhon.jairo.vallejo@gmail.com", "Inspeccionar esquema del dataset valledata en BigQuery"],
            ["ALT-LOD-003", "Carga (Bronze)", "Detección de respuesta status=SIMULATED en Composer", "CRÍTICO", "Todos los DAGs src_load_*", "jhon.jairo.vallejo@gmail.com", "Validar conexión nativa google_cloud_default en Composer"],
            ["ALT-LOD-004", "Carga (Bronze)", "Incompatibilidad de tipos de datos o corrupción en parseo de CSV/Excel", "ALTO", "src_load_sipsa, src_load_crops, src_load_ckan_comentarios", "jhon.jairo.vallejo@gmail.com", "Revisar cambios de formato en fuentes raw o parseo pandas"],
            ["ALT-TRF-001", "Transformación (Silver)", "Error en modelo de Clasificación de Sentimiento NLP (pysentimiento/VADER)", "MEDIO", "src_transform_ckan_comentarios", "jhon.jairo.vallejo@gmail.com", "Verificar dependencias NLP y memoria de Workers"],
            ["ALT-TRF-002", "Transformación (Gold)", "Fallo en cruce de llaves (Año, Municipio, Cultivo, Semestre) en Consolidado", "CRÍTICO", "src_transform_consolidado, src_transform_crops", "jhon.jairo.vallejo@gmail.com", "Validar integridad de tablas Bronze stg_*"],
            ["ALT-TRF-003", "Transformación (Gold)", "Generación de dataset consolidado con 0 registros (Data Quality Anomaly)", "CRÍTICO", "src_transform_consolidado", "jhon.jairo.vallejo@gmail.com", "Auditar coincidencia de rangos de años entre SIPSA, ONI y Cultivos"],
            ["ALT-TRF-004", "Transformación GIS (Gold)", "Fallo en generación de polígonos WKT/GeoJSON, centroides WGS84 o Haversine Cavasa", "CRÍTICO", "src_transform_spatial", "jhon.jairo.vallejo@gmail.com", "Verificar coordenadas municipales y tabla BigQuery Gold gold_cultivos_valle_geo"],
            ["ALT-ORC-001", "Orquestación Maestro", "Fallo o timeout en disparo/espera de DAGs hijos vía TriggerDagRunOperator", "CRÍTICO", "src_orchestrator_master", "jhon.jairo.vallejo@gmail.com", "Revisar logs del DAG hijo que falló y estado del scheduler Airflow"],
            ["ALT-SYS-001", "Infraestructura", "Excepción de runtime no capturada en ejecución de tareas de Python", "CRÍTICO", "Todos los 15 DAGs", "jhon.jairo.vallejo@gmail.com", "Revisar traceback detallado en logs de TaskInstance"],
            ["ALT-SYS-002", "Infraestructura", "Agotamiento definitivo de reintentos (try_number > retries)", "CRÍTICO", "Todos los 15 DAGs", "jhon.jairo.vallejo@gmail.com", "Intervención manual de soporte y revisión de recursos de Airflow"],
            ["ALT-SYS-003", "Infraestructura", "Fallo al enviar notificación por correo SMTP (send_email)", "BAJO", "Callback airflow_failure_alarm", "jhon.jairo.vallejo@gmail.com", "Verificar servidor SMTP en airflow.cfg / variables de entorno Composer"],
            ["ALT-SYS-004", "Infraestructura", "Incompatibilidad de SDK Airflow (Composer 2 vs Composer 3 import SDK)", "MEDIO", "Todos los 15 DAGs", "jhon.jairo.vallejo@gmail.com", "Asegurar bloque try... import airflow.sdk ... except... import airflow.decorators"],
        ]

        while len(t2.rows) < len(alarms_data):
            t2.add_row()

        for r_i, row_values in enumerate(alarms_data):
            row = t2.rows[r_i]
            for c_i, val in enumerate(row_values):
                if c_i < len(row.cells):
                    bg = "003366" if r_i == 0 else ("F8F9FA" if r_i % 2 == 1 else "FFFFFF")
                    tc = RGBColor(255, 255, 255) if r_i == 0 else RGBColor(0, 0, 0)
                    set_cell(row.cells[c_i], val, font_name="Arial Narrow", font_size=Pt(9), bold=(r_i == 0), bg_hex=bg, text_color=tc)

    # 5. Clear old body paragraphs starting at P70 and write Alarms document sections
    p_start = 70
    total_p = len(doc.paragraphs)
    for i in range(p_start, total_p):
        doc.paragraphs[i].text = ""

    current_idx = p_start

    def add_p(text, style="normal", font_name="Arial Narrow", font_size=Pt(11), bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        nonlocal current_idx
        if current_idx < len(doc.paragraphs):
            p = doc.paragraphs[current_idx]
            p.style = doc.styles[style] if style in doc.styles else doc.styles["normal"]
        else:
            p = doc.add_paragraph(style=style)
        current_idx += 1
        p.alignment = align
        run = p.add_run(text)
        run.font.name = font_name
        if font_size:
            run.font.size = font_size
        run.bold = bold
        run.italic = italic
        return p

    # Section 1
    add_p("1. Visión General del Sistema de Alarmas en Producción", style="Heading 1", font_name="Arial", font_size=Pt(14), bold=True)
    add_p("El sistema de alarmas de la plataforma ValleDATA provee una arquitectura centralizada y resiliente para detectar, clasificar, notificar y registrar cualquier fallo u anomalía operativa durante la ejecución de los 15 DAGs del pipeline Medallion y Orquestador Maestro en Google Cloud Composer (proyecto GCP datagov-477214, entorno composer-gdv).")

    # Section 2
    add_p("2. Configuración Global de Alarmas y Destinatario de Correo", style="Heading 1", font_name="Arial", font_size=Pt(14), bold=True)
    add_p("El comportamiento de las alarmas se parametriza globalmente en config.yaml bajo la clave airflow_alerts apuntando al destinatario oficial:")
    add_p("•  Destinatario Oficial: jhon.jairo.vallejo@gmail.com")
    add_p("•  Notificación en Fallo: email_on_failure: true")
    add_p("•  Reintentos Automáticos: 1 intento con retrazo de 5 minutos (retry_delay_minutes: 5)")
    add_p("•  Estados de Disparo: ERROR, PARTIAL, FAILED")

    # Section 3
    add_p("3. Matriz Consolidada de Eventos de Alarma", style="Heading 1", font_name="Arial", font_size=Pt(14), bold=True)

    # Section 4
    add_p("4. Procedimientos de Respuesta e Incidentes", style="Heading 1", font_name="Arial", font_size=Pt(14), bold=True)
    add_p("Ante la recepción de una alerta enviada a jhon.jairo.vallejo@gmail.com, el equipo de ingeniería debe seguir la siguiente secuencia:")
    add_p("1. Verificar el log de la tarea mediante el enlace incluido en el correo (Composer Airflow UI).")
    add_p("2. Validar la disponibilidad del portal origen (DANE, NOAA, Gobernación, datos.gov.co o CKAN).")
    add_p("3. Si el fallo es por latencia de red o timeout de API, activar la re-ejecución manual en Airflow.")
    add_p("4. Si el error es por esquema de BigQuery o GIS espacial, corregir el tipo de dato en la función de transformación.")

    # Section 5
    add_p("5. Firma y Validez del Documento", style="Heading 1", font_name="Arial", font_size=Pt(14), bold=True)
    add_p("Este documento constituye la matriz oficial de eventos y alarmas para auditorías de software y continuidad operativa de la plataforma ValleDATA.")
    add_p("FIN DEL DOCUMENTO.", align=WD_ALIGN_PARAGRAPH.CENTER)

    targets = [
        Path("/home/jjvallej/work/enigma/pyimport/15.1. MATRIZ_DE_EVENTOS_ALARMAS_DAGS.docx"),
        Path("/home/jjvallej/work/enigma/pyimport/15.1 VF - MATRIZ_DE_EVENTOS_ALARMAS_DAGS.docx"),
        Path("/home/jjvallej/work/enigma/pyimport/1.3.1.4 MATRIZ_DE_EVENTOS_ALARMAS_DAGS.docx"),
        Path("/home/jjvallej/work/enigma/pyimport/MATRIZ_DE_EVENTOS_ALARMAS_DAGS.docx"),
        Path("/home/jjvallej/work/enigma/airflow/15.1. MATRIZ_DE_EVENTOS_ALARMAS_DAGS.docx"),
        Path("/home/jjvallej/work/enigma/airflow/15.1 VF - MATRIZ_DE_EVENTOS_ALARMAS_DAGS.docx"),
        Path("/home/jjvallej/work/enigma/airflow/1.3.1.4 MATRIZ_DE_EVENTOS_ALARMAS_DAGS.docx"),
        Path("/home/jjvallej/work/enigma/airflow/MATRIZ_DE_EVENTOS_ALARMAS_DAGS.docx"),
    ]

    for target in targets:
        target.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(target))
        print(f"✅ Matriz de alarmas generada exitosamente en: {target}")

if __name__ == "__main__":
    main()
