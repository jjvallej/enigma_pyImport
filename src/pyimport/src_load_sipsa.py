"""ETAPA LOAD SIPSA (src_load_sipsa.py)
Consolida anexos SIPSA (hoja "Anexo 1" / "1"):
  - Columna A = cultivo/alimento
  - Columna J = precio Cali (o la columna cuyo encabezado sea "Cali")
  - Año/mes desde el nombre del archivo
Calcula el promedio anual por cultivo y carga:
  datagov-477214.valledata.bronze_agri_sipsa
"""

from __future__ import annotations

from collections import defaultdict
import csv
from dataclasses import dataclass
from pathlib import Path
import re
import sys
from typing import Any, Dict, Iterable

CURRENT_DIR = str(Path(__file__).resolve().parent)
if CURRENT_DIR not in sys.path:
    sys.path.insert(0, CURRENT_DIR)

from src_common import (  # noqa: E402
    ensure_xlrd,
    get_bigquery_client,
    get_composer_params,
    get_connection_id,
    get_raw_root,
    list_storage,
    load_config,
    materialize_local,
    require_config_value,
    storage_join,
    write_bytes,
    get_airflow_dag_kwargs,
    run_with_airflow_alarm,
)

MONTH_ABBR = {
    "ene": "01", "feb": "02", "mar": "03", "abr": "04", "may": "05", "jun": "06",
    "jul": "07", "ago": "08", "sep": "09", "oct": "10", "nov": "11", "dic": "12",
}
MONTH_TOKEN = "|".join(MONTH_ABBR)

FILENAME_PATTERNS: tuple[re.Pattern[str], ...] = (
    re.compile(rf"^anex_mensual_(?P<mes>{MONTH_TOKEN})_(?P<anio>\d{{4}})\.(?:xls|xlsx)$", re.IGNORECASE),
    re.compile(rf"^anexo_mensual_SIPSA_mayoristas_(?P<mes>{MONTH_TOKEN})_(?P<anio>\d{{4}})\.xlsx$", re.IGNORECASE),
    re.compile(rf"^anex-SIPSAMensual-(?P<mes>{MONTH_TOKEN})(?P<anio>\d{{4}})\.xlsx$", re.IGNORECASE),
)

SKIP_NAMES = {
    "hortalizas y verduras", "frutas frescas", "tubérculos y plátanos", "tuberculos y platanos",
    "granos, cárnicos y procesados", "granos, carnicos y procesados", "producto", "precio $/kg",
    "fuente: dane", "precio", "var %", "var%", "cali",
}

SKIP_PREFIXES = (
    "anexo ", "variación", "variacion", "comportamiento", "sistema de información",
    "sistema de informacion", "var%", "n.d.", "fuente:", "*variedad", "-:",
)

# Columna J (1-based = 10, 0-based = 9) = Cali por defecto.
DEFAULT_CALI_COL = 9


@dataclass(frozen=True)
class MonthlyPrice:
    anio: str
    mes: str
    cultivo: str
    precio_cali: float


@dataclass(frozen=True)
class YearlyPrice:
    anio: str
    cultivo: str
    precio_promedio_cali: float
    meses_con_dato: int


def parse_period_from_filename(filename: str) -> tuple[int, int]:
    name = Path(filename).name
    for pattern in FILENAME_PATTERNS:
        match = pattern.match(name)
        if match:
            year = int(match.group("anio"))
            month = int(MONTH_ABBR[match.group("mes").lower()])
            return year, month
    raise ValueError(f"Nombre de archivo no reconocido: {filename}")


def _normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    text = re.sub(r"[*]+", "", str(value)).strip()
    if text.lower() in {"none", "nan"}:
        return ""
    return text


def _is_cultivo_name(name: str) -> bool:
    if not name:
        return False
    lower = name.lower()
    if lower in SKIP_NAMES:
        return False
    if any(lower.startswith(prefix) for prefix in SKIP_PREFIXES):
        return False
    if re.match(r"^[\d\s.,:\-*]+$", lower):
        return False
    return True


def _parse_price(value: Any) -> float | None:
    text = _normalize_cell(value)
    if not text:
        return None
    lower = text.lower()
    if lower in {"n.d.", "-", "precio", "var %", "var%", "cali"}:
        return None
    try:
        return float(text.replace(",", "."))
    except ValueError:
        return None


def _pick_anexo1_sheet(sheet_names: list[str]) -> str:
    """Prioriza hoja 'Anexo 1' o '1' (precios), nunca la de variaciones."""
    normalized = {name.strip().lower(): name for name in sheet_names}
    for key in ("anexo 1", "1"):
        if key in normalized:
            return normalized[key]
    for name in sheet_names:
        if name.strip().lower() not in {"indice", "índice", "index", "anexo 2", "2", "anexo 3", "3"}:
            return name
    return sheet_names[0]


def _find_cali_col(matrix: list[list[Any]], default: int = DEFAULT_CALI_COL) -> int:
    """Busca la columna cuyo encabezado sea 'Cali' (normalmente J)."""
    for row in matrix[:40]:
        for idx, cell in enumerate(row):
            if _normalize_cell(cell).lower() == "cali":
                return idx
    return default


def _read_sheet_matrix(path: Path, sheet_name: str) -> list[list[Any]]:
    """Lee la hoja. .xls usa xlrd (sistema o src_xlrd_vendor.zip); .xlsx usa pandas/openpyxl."""
    suffix = path.suffix.lower()

    if suffix == ".xls":
        xlrd = ensure_xlrd()
        wb = xlrd.open_workbook(path)
        sheet = wb.sheet_by_name(sheet_name)
        return [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]

    try:
        import pandas as pd

        df = pd.read_excel(path, sheet_name=sheet_name, header=None, dtype=object)
        return df.where(pd.notnull(df), None).values.tolist()
    except Exception as pd_exc:
        print(f"ℹ️ pandas no pudo leer {path.name}/{sheet_name}: {pd_exc}", flush=True)

    if suffix == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            ws = wb[sheet_name]
            return [list(row) for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()

    raise ValueError(f"No se pudo leer {path}")


def _list_sheet_names(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".xls":
        xlrd = ensure_xlrd()
        return list(xlrd.open_workbook(path).sheet_names())

    try:
        import pandas as pd

        xl = pd.ExcelFile(path)
        return list(xl.sheet_names)
    except Exception:
        pass

    if suffix == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
    raise ValueError(f"Extensión no soportada: {path.suffix}")


def extract_monthly_prices(
    file_path: Path,
    source_name: str | None = None,
) -> list[MonthlyPrice]:
    """Extrae cultivo (A) y precio Cali (J) de la hoja Anexo 1.

    source_name: nombre original del objeto (GCS/local). Si se omite, usa file_path.name.
    """
    name_for_period = source_name or file_path.name
    try:
        year_val, month_val = parse_period_from_filename(name_for_period)
    except ValueError:
        print(f"⚠️ Nombre no reconocido, se omite: {name_for_period}", flush=True)
        return []

    anio = str(year_val)
    mes = f"{month_val:02d}"

    try:
        sheets = _list_sheet_names(file_path)
        sheet_name = _pick_anexo1_sheet(sheets)
        matrix = _read_sheet_matrix(file_path, sheet_name)
    except Exception as exc:
        print(f"⚠️ Error leyendo Excel {file_path.name}: {exc}", flush=True)
        return []

    if not matrix:
        return []

    cali_col = _find_cali_col(matrix)
    print(
        f"   ↳ hoja={sheet_name!r} col_cali={cali_col + 1} (A=1,J=10) período={anio}-{mes}",
        flush=True,
    )

    records: list[MonthlyPrice] = []
    for row in matrix:
        if not row:
            continue
        cultivo = _normalize_cell(row[0] if len(row) > 0 else "")
        if not _is_cultivo_name(cultivo):
            continue
        precio = _parse_price(row[cali_col] if len(row) > cali_col else None)
        if precio is None:
            continue
        records.append(MonthlyPrice(anio=anio, mes=mes, cultivo=cultivo, precio_cali=precio))
    return records


def average_by_year(rows: Iterable[MonthlyPrice]) -> list[YearlyPrice]:
    """Promedio de precios Cali por año y cultivo."""
    buckets: dict[tuple[str, str], list[float]] = defaultdict(list)
    for row in rows:
        buckets[(row.anio, row.cultivo)].append(row.precio_cali)

    result: list[YearlyPrice] = []
    for (anio, cultivo), prices in sorted(buckets.items(), key=lambda x: (x[0][0], x[0][1].lower())):
        result.append(
            YearlyPrice(
                anio=anio,
                cultivo=cultivo,
                precio_promedio_cali=round(sum(prices) / len(prices), 2),
                meses_con_dato=len(prices),
            )
        )
    return result


def write_yearly_csv(rows: Iterable[YearlyPrice], output_path: Path) -> int:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with output_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["anio", "cultivo", "precio_promedio_cali", "meses_con_dato"])
        for r in rows:
            writer.writerow([r.anio, r.cultivo, f"{r.precio_promedio_cali:.2f}", r.meses_con_dato])
            count += 1
    return count


def list_sipsa_anexos(download_root: str, cfg: Dict[str, Any] | None = None) -> list[str]:
    """Lista anexos SIPSA (anex* / anexo*)."""
    uris = list_storage(download_root, name_prefix="anex", cfg=cfg)
    # Incluir también si el listado local/GCS devolvió mixto
    return [u for u in uris if Path(u).suffix.lower() in {".xls", ".xlsx"}]


def import_sipsa_to_csv(
    download_root: str,
    output_path: Path,
    cfg: Dict[str, Any] | None = None,
) -> dict[str, object]:
    uris = list_sipsa_anexos(download_root, cfg=cfg)
    if not uris:
        raise FileNotFoundError(
            f"No hay anexos SIPSA (.xls/.xlsx) en {download_root}. Ejecute primero src_ingest_sipsa."
        )

    monthly: list[MonthlyPrice] = []
    files_ok = 0
    print(f"📂 [SRC_LOAD_SIPSA] Leyendo {len(uris)} anexos desde {download_root}", flush=True)

    for i, uri in enumerate(uris, start=1):
        original_name = Path(uri).name
        local = materialize_local(uri, cfg=cfg)
        rows = extract_monthly_prices(local, source_name=original_name)
        print(f"   [{i}/{len(uris)}] {original_name}: {len(rows)} precios mensuales", flush=True)
        if rows:
            files_ok += 1
            monthly.extend(rows)

    if not monthly:
        raise RuntimeError(
            "No se extrajo ningún precio Cali de los anexos. "
            "Verifique hoja 'Anexo 1'/'1', columna J y dependencias (pandas/openpyxl/xlrd)."
        )

    yearly = average_by_year(monthly)
    written = write_yearly_csv(yearly, output_path)
    years = sorted({r.anio for r in yearly})
    print(
        f"🧾 [SRC_LOAD_SIPSA] Promedio anual: {written} filas "
        f"(mensuales={len(monthly)}, archivos_ok={files_ok}/{len(uris)}, "
        f"años={years[0]}…{years[-1]}) -> {output_path}",
        flush=True,
    )
    return {
        "files_found": len(uris),
        "files_processed": files_ok,
        "monthly_rows": len(monthly),
        "rows": written,
        "years": years,
        "output": str(output_path),
    }


def run_load_sipsa(config: Dict[str, Any] | None = None) -> Dict[str, Any]:
    cfg = config or load_config()
    paths_cfg = require_config_value(cfg, "paths")
    bq_cfg = require_config_value(cfg, "bigquery")
    google_cloud_default = get_connection_id(cfg, "google_cloud_default")
    composer = get_composer_params(cfg)

    raw_root = get_raw_root(cfg)
    output_csv = Path(require_config_value(paths_cfg, "sipsa_csv"))

    project_id = require_config_value(bq_cfg, "project_id")
    dataset_bronze = require_config_value(bq_cfg, "datasets", "bronze")
    table_bronze = require_config_value(bq_cfg, "tables", "sipsa_bronze")
    table_ref = f"{project_id}.{dataset_bronze}.{table_bronze}"

    print(
        f"📦 [SRC_LOAD_SIPSA] raw={raw_root} -> CSV={output_csv} -> BQ={table_ref} | "
        f"Composer={composer['environment']} ({composer['location']}) | conn={google_cloud_default}",
        flush=True,
    )
    res = import_sipsa_to_csv(download_root=raw_root, output_path=output_csv, cfg=cfg)

    try:
        staging_uri = storage_join(raw_root, output_csv.name)
        write_bytes(staging_uri, output_csv.read_bytes(), cfg=cfg)
        res["staging_uri"] = staging_uri
        print(f"☁️  [SRC_LOAD_SIPSA] CSV en staging: {staging_uri}", flush=True)
    except Exception as exc:
        print(f"ℹ️ No se pudo copiar CSV a staging: {exc}", flush=True)

    from google.cloud import bigquery

    client = get_bigquery_client(cfg, project_id, require_config_value(bq_cfg, "location"))
    job_config = bigquery.LoadJobConfig(
        source_format=bigquery.SourceFormat.CSV,
        skip_leading_rows=1,
        write_disposition=bigquery.WriteDisposition.WRITE_TRUNCATE,
        schema=[
            bigquery.SchemaField("anio", "STRING"),
            bigquery.SchemaField("cultivo", "STRING"),
            bigquery.SchemaField("precio_promedio_cali", "FLOAT"),
            bigquery.SchemaField("meses_con_dato", "INTEGER"),
        ],
    )
    with output_csv.open("rb") as sf:
        job = client.load_table_from_file(sf, table_ref, job_config=job_config)
    job.result()

    table = client.get_table(table_ref)
    res["bigquery_table"] = table_ref
    res["bigquery_rows"] = table.num_rows
    res["status"] = "SUCCESS"
    print(
        f"✅ [SRC_LOAD_SIPSA] Cargado {table.num_rows} filas en {table_ref}",
        flush=True,
    )
    if not table.num_rows:
        raise RuntimeError(f"La tabla {table_ref} quedó vacía tras la carga.")
    return res


if __name__ == "__main__":
    run_load_sipsa()

try:
    from datetime import datetime
    from airflow.decorators import dag, task

    from src_common import get_airflow_dag_kwargs, run_with_airflow_alarm

    @dag(
        dag_id="src_load_sipsa",
        description="Consolida SIPSA Anexo1 (precio Cali promedio anual) -> bronze_agri_sipsa",
        start_date=datetime(2000, 1, 1),
        schedule=None,
        catchup=False,
        tags=["valledata", "sipsa", "load", "bronze"],
        **get_airflow_dag_kwargs(),
    )
    def load_sipsa_dag():
        @task(task_id="run_load_sipsa")
        def execute_load() -> dict[str, object]:
            return run_with_airflow_alarm(run_load_sipsa)

        execute_load()

    dag_single = load_sipsa_dag()
except ImportError:
    pass
